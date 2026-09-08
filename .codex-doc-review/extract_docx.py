from docx import Document

path = r"C:\Users\ASUS\WPSDrive\1625175705\WPS企业云盘\北京建筑大学\我的企业文档\技术相关知识.docx"
doc = Document(path)
print(f"PARAGRAPHS {len(doc.paragraphs)} TABLES {len(doc.tables)}")
for i, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip():
        print(f"{i}: [{paragraph.style.name}] {paragraph.text}")
for table_index, table in enumerate(doc.tables):
    print(f"\nTABLE {table_index}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
